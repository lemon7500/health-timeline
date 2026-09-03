from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "dist" / "病程日历-客户使用指南.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
PALE_YELLOW = "FFF4CE"
PALE_GREEN = "E8F5E9"
TEXT = RGBColor(38, 45, 52)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, keep: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = keep


def set_run_font(run, size=None, bold=None, color=None, east_asia="Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_field(paragraph) -> None:
    paragraph.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    paragraph.add_run(" 页")


def add_callout(doc: Document, title: str, body: str, fill=PALE_YELLOW) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.35)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, 10.5, True, DARK_BLUE)
    r = p.add_run("\n" + body)
    set_run_font(r, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def create_numbering_instance(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 1
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(tabs)
    p_pr.append(ind)
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.insert(0, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_steps(doc: Document, items: list[str]) -> None:
    num_id = create_numbering_instance(doc)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_node])
        p_pr.append(num_pr)
        p.add_run(item)


def add_heading(doc: Document, text: str, level=1) -> None:
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)


def add_body(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_kv_table(doc: Document, rows: list[tuple[str, str]], widths=(1.45, 4.9)) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(widths[0])
        cells[1].width = Inches(widths[1])
        set_cell_shading(cells[0], PALE_BLUE)
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cells[0].paragraphs[0].add_run(key)
        set_run_font(r, 10, True, DARK_BLUE)
        r = cells[1].paragraphs[0].add_run(value)
        set_run_font(r, 10)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    settings = {
        "Title": (26, DARK_BLUE, 0, 8),
        "Subtitle": (12, "657786", 0, 14),
        "Heading 1": (16, BLUE, 16, 7),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (11.5, DARK_BLUE, 9, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.38)
        style.paragraph_format.first_line_indent = Inches(-0.19)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.18


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    p = first_header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("客户资料 · 请妥善保管")
    set_run_font(r, 8.5, False, "7A8793")

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("病程日历 1.2.0  |  客户使用指南")
    set_run_font(r, 8.5, False, "657786")
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), PALE_BLUE)
    borders.append(bottom)
    p_pr.append(borders)

    for footer in (section.footer, section.first_page_footer):
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        add_page_field(p)
        for run in p.runs:
            set_run_font(run, 8.5, False, "7A8793")


def build() -> None:
    doc = Document()
    style_document(doc)
    add_header_footer(doc)

    p = doc.add_paragraph(style="Title")
    p.add_run("病程日历使用指南")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("个人病情记录、复查提醒与用药管理")

    add_kv_table(
        doc,
        [
            ("适用版本", "1.2.0"),
            ("适用设备", "Android 8.0 及以上的华为、荣耀、Redmi/小米及其他 Android 手机"),
            ("不支持", "HarmonyOS NEXT 原生系统"),
            ("数据方式", "完全离线，本地加密存储"),
        ],
    )
    doc.add_paragraph()
    add_callout(
        doc,
        "使用前请先阅读",
        "数据只保存在手机本地。卸载应用、清除应用数据或手机损坏，都可能导致本机数据消失。请定期导出加密备份，并把备份文件保存到另一台设备、U 盘或电脑。已安装旧版本时，请直接覆盖安装新版，不要先卸载。",
    )

    add_heading(doc, "1. 软件用途", 1)
    add_body(doc, "“病程日历”用于个人病情资料整理和提醒，可帮助您：")
    add_bullets(
        doc,
        [
            "为本人和家人分别建立档案，并在各页面切换当前成员。",
            "按日历查看每次就诊、病情变化、诊断、治疗和备注。",
            "保存每次检查报告的多张图片和多个 PDF。",
            "设置定期复查提醒，例如“每 3 个月一次”或“每周四”。",
            "登记药物疗程、每日服药时间，并记录“已服”或“跳过”。",
            "使用系统指纹、面容或锁屏密码保护应用。",
            "导出全家庭密码加密备份，在换机或重新安装后恢复。",
        ],
    )
    add_callout(doc, "医疗提示", "本应用只用于记录和提醒，不能替代医生诊断、处方或用药指导。", PALE_GREEN)

    add_heading(doc, "2. 安装应用", 1)
    add_steps(
        doc,
        [
            "从 GitHub Releases 下载“HealthTimeline-1.2.0.apk”并保存到手机。",
            "在“文件管理”中点击 APK。",
            "首次安装时，按系统提示允许当前文件管理器“安装未知应用”。",
            "安装完成后，可关闭“安装未知应用”权限。",
            "打开“病程日历”，按提示允许通知。",
        ],
    )
    add_body(doc, "如果提示“无法安装”，请确认手机为 Android 8.0 或以上、APK 文件下载完整，并且不是 HarmonyOS NEXT 原生系统。")

    add_heading(doc, "3. 首次使用必须检查的权限", 1)
    add_body(doc, "为了让复查和服药提醒尽量准时，请完成以下设置：")
    add_steps(
        doc,
        [
            "通知权限：允许“病程日历”发送通知。",
            "精确闹钟：在应用“设置”页按提示允许精确提醒；拒绝后仍可提醒，但可能延迟。",
            "后台运行：允许应用后台活动和自动启动。",
            "电池管理：将应用设为“不限制”或“不优化”。",
        ],
    )
    add_heading(doc, "华为/荣耀", 2)
    add_body(doc, "通常可在“设置 → 应用和服务 → 应用启动管理”中找到本应用，关闭自动管理后允许“自动启动、关联启动、后台活动”。")
    add_heading(doc, "Redmi/小米", 2)
    add_body(doc, "通常可在“设置 → 应用设置 → 应用管理 → 病程日历”中打开“自启动”，并在“省电策略”中选择“无限制”。不同机型和系统版本的菜单名称可能略有不同。")

    doc.add_page_break()
    add_heading(doc, "4. 四个主要页面", 1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(5.1)
    hdr = table.rows[0].cells
    hdr[0].text = "页面"
    hdr[1].text = "用途"
    set_repeat_table_header(table.rows[0])
    for cell in hdr:
        set_cell_shading(cell, PALE_BLUE)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, 10, True, DARK_BLUE)
    for name, purpose in [
        ("日历", "切换家庭成员；查看、搜索、新增和修改病历"),
        ("复查", "切换成员；设置一次性或周期性复查，处理待复查事项"),
        ("用药", "切换成员；登记药物、疗程和每日服药记录"),
        ("设置", "管理家庭档案和应用锁，查看权限，导出或恢复全家庭备份"),
    ]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = purpose
        for cell in cells:
            set_cell_margins(cell)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, 10)

    add_heading(doc, "5. 新增病情或就诊记录", 1)
    add_steps(
        doc,
        [
            "打开“日历”，切换到需要记录的月份。",
            "选择日期，点击新增按钮。",
            "选择或新建病情分类，例如“乳腺复查”或“上颌窦炎术后”。",
            "填写标题，并按需要填写就诊阶段、症状、诊断、治疗、用药、医院、医生和备注。",
            "点击“保存”。",
        ],
    )
    add_body(doc, "日历中每个日期最多显示两个标题；当天记录更多时会显示“+N”。点击日期即可查看全部记录。已有历史记录引用的病情分类只能归档，不能直接删除。")

    add_heading(doc, "6. 保存和查看检查报告", 1)
    add_steps(
        doc,
        [
            "打开一条病历记录。",
            "在附件区域选择拍照、选择图片或选择 PDF。",
            "选中的文件会复制到应用的私有存储中，原文件移动或改名不影响应用内副本。",
            "保存后，点击图片可缩放查看；点击 PDF 可逐页查看。",
        ],
    )
    add_callout(doc, "附件限制", "每条病历可保存多张图片和多个 PDF。单个文件不能超过 100 MB。文件损坏、格式不支持或手机空间不足时，应用会取消本次导入，不会保留半份附件。", PALE_BLUE)

    add_heading(doc, "7. 设置复查提醒", 1)
    add_steps(
        doc,
        [
            "打开“复查”，点击新增。",
            "填写复查名称、首次日期、提醒时间和提前天数。",
            "选择重复方式：一次性、每 N 天、每 N 周指定星期，或每 N 月指定日期。",
            "保存后，系统会安排下一次提醒。",
            "完成复查后标记“已完成”；不需要本次复查时可标记“已跳过”。",
        ],
    )
    add_heading(doc, "计算示例", 2)
    add_bullets(
        doc,
        [
            "7 月 9 日设置“每 3 个月”，下一次为 10 月 9 日。",
            "设置“每周四”，应用会自动跨月、跨年计算下一个周四。",
            "1 月 31 日设置每月提醒，2 月使用当月最后一天，3 月仍回到 31 日。",
        ],
    )
    add_body(doc, "未处理的复查会继续显示为待办或逾期。只有标记“已完成”或“已跳过”后，周期提醒才推进到下一次。")

    add_heading(doc, "8. 登记用药和每日打卡", 1)
    add_steps(
        doc,
        [
            "打开“用药”，点击新增药物。",
            "填写药名、剂量、单位、开始日期、结束日期和服用说明。",
            "选择“按需服用”，或设置每天一个或多个服药时间，例如 08:00、20:00。",
            "在“今日用药”中选择“已服”或“跳过”；也可在通知中快捷标记“已服”。",
            "修改药物剂量只影响之后的计划，过去的用药记录会保留当时的剂量。",
        ],
    )
    add_body(doc, "未点击“已服”或“跳过”的项目保持“未记录”，应用不会自动认定为漏服。疗程结束后不再产生新的提醒，历史记录仍然保留。")

    add_heading(doc, "9. 导出加密备份", 1)
    add_callout(doc, "建议频率", "至少每周备份一次，并在添加重要检查报告、换手机或更新应用前额外备份。", PALE_GREEN)
    add_steps(
        doc,
        [
            "打开“设置 → 导出加密备份”。",
            "输入至少 8 位密码并再次确认。",
            "选择保存位置，生成 .htbackup 文件。",
            "将备份文件复制到电脑、U 盘或另一台可靠设备。",
            "备份文件和密码请分开保管。",
        ],
    )
    add_callout(doc, "密码无法找回", "v3 备份包含全部家庭成员的病历、复查、用药记录和附件。密码不会保存在应用中，忘记后无法恢复该备份。", PALE_YELLOW)

    add_heading(doc, "10. 恢复备份或更换手机", 1)
    add_steps(
        doc,
        [
            "在新手机安装“病程日历”。",
            "将 .htbackup 备份文件复制到新手机。",
            "打开“设置 → 从备份恢复”，选择备份文件。",
            "输入导出时设置的密码并确认。",
            "应用会先检查密码、版本、成员引用、文件完整性和可用空间，再显示安全合并或整体替换选项。",
        ],
    )
    add_callout(doc, "整体替换前先备份", "整体替换会删除本机独有资料，因此应用会先强制导出当前数据安全备份。导出或验证失败时，现有数据不会被修改。", PALE_BLUE)

    add_heading(doc, "11. 更新应用", 1)
    add_body(doc, "如果手机已经安装旧版本，请从 GitHub 下载新版 APK 后直接覆盖安装。不要先卸载旧版，也不要清除应用数据。1.1.0 升级到 1.2.0 后原资料会归入“本人”；如果系统提示签名不一致，请停止操作并联系软件提供方。")

    add_heading(doc, "12. 常见问题", 1)
    questions = [
        ("收不到提醒或提醒延迟", "检查通知权限、精确闹钟、自动启动、后台活动和电池“不限制”设置。系统省电模式可能延迟通知；修改设置后，可重新打开应用让提醒重新安排。"),
        ("Redmi/小米手机锁屏后不提醒", "打开本应用的“自启动”，将省电策略设为“无限制”，并确认通知允许显示。HyperOS/MIUI 更新后，建议再次检查这些设置。"),
        ("华为/荣耀手机锁屏后不提醒", "在“应用启动管理”中改为手动管理，并允许自动启动、关联启动和后台活动；同时关闭对本应用的电池优化。"),
        ("报告无法导入", "请确认文件是常见图片或 PDF、单个文件不超过 100 MB、文件没有损坏，并且手机剩余空间充足。"),
        ("忘记备份密码", "加密备份密码无法找回。若旧手机中的应用数据仍在，请重新导出备份并设置新密码。"),
        ("卸载后记录不见了", "卸载应用会删除其私有数据。重新安装后只能通过之前导出的 .htbackup 文件恢复。"),
    ]
    for question, answer in questions:
        add_heading(doc, question, 2)
        add_body(doc, answer)

    add_heading(doc, "13. 隐私与安全说明", 1)
    add_bullets(
        doc,
        [
            "应用不申请网络权限，不登录账号，不上传病历，也不进行云同步。",
            "数据库经过加密，图片和 PDF 保存在应用私有目录。",
            "可开启系统身份验证；锁屏通知隐藏成员姓名、病情和药名。",
            "手机遗失、损坏、恢复出厂设置、卸载或清除数据仍可能造成数据丢失；离线应用无法替您找回未备份的数据。",
        ],
    )

    doc.add_page_break()
    add_heading(doc, "14. 交付客户前快速检查", 1)
    add_steps(
        doc,
        [
            "能正常打开应用并新增一条测试病历。",
            "已允许通知、精确闹钟和后台运行。",
            "已设置一条近期测试提醒并确认能够收到。",
            "已成功导出一次加密备份，并确认记得密码。",
            "已把备份复制到手机之外的安全位置。",
        ],
    )
    add_callout(doc, "需要技术支持时", "请联系软件提供方，并说明手机品牌、型号、系统版本和遇到问题的具体步骤。", PALE_GREEN)

    core_props = doc.core_properties
    core_props.title = "病程日历使用指南"
    core_props.subject = "病程日历 Android 应用客户使用说明"
    core_props.author = "病程日历"
    core_props.keywords = "病程日历, Android, 使用指南, 复查提醒, 用药记录, 加密备份"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
